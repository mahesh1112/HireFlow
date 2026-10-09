import os
import uuid
import aiofiles
from pathlib import Path
from fastapi import UploadFile, HTTPException, status
from app.config import settings


async def save_upload_file(file: UploadFile, user_id: str) -> tuple[str, str]:
    """
    Save an uploaded file to disk. Returns (file_path, content_type).
    """
    if file.content_type not in settings.ALLOWED_FILE_TYPES:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"File type '{file.content_type}' is not allowed. Only PDF and DOCX are accepted.",
        )

    # Read file content and check size
    content = await file.read()
    if len(content) > settings.MAX_FILE_SIZE:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail="File exceeds maximum size of 10MB",
        )

    # Ensure upload directory exists
    upload_dir = Path(settings.UPLOAD_DIR) / user_id
    upload_dir.mkdir(parents=True, exist_ok=True)

    # Generate unique filename
    ext = Path(file.filename or "resume").suffix or ".pdf"
    unique_name = f"{uuid.uuid4()}{ext}"
    file_path = upload_dir / unique_name

    # Write to disk
    async with aiofiles.open(file_path, "wb") as f:
        await f.write(content)

    return str(file_path), file.content_type or "application/pdf"


def extract_text_from_pdf(file_path: str) -> str:
    """Extract plain text from a PDF file using pdfplumber."""
    try:
        import pdfplumber

        with pdfplumber.open(file_path) as pdf:
            pages = [page.extract_text() or "" for page in pdf.pages]
        return "\n".join(pages)
    except Exception as exc:
        raise RuntimeError(f"Failed to extract text from PDF: {exc}") from exc


def extract_text_from_docx(file_path: str) -> str:
    """Extract plain text from a DOCX file."""
    try:
        from docx import Document

        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs]
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.append(cell.text)
        return "\n".join(paragraphs)
    except Exception as exc:
        raise RuntimeError(f"Failed to extract text from DOCX: {exc}") from exc


def extract_text(file_path: str, content_type: str) -> str:
    """Route to the correct extractor based on content type."""
    if content_type == "application/pdf":
        return extract_text_from_pdf(file_path)
    elif content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported content type: {content_type}")


def cleanup_file(file_path: str) -> None:
    """Remove a temporary file."""
    try:
        os.remove(file_path)
    except OSError:
        pass


# ─── BYTEA-based functions (Gap 11.14) ────────────────────────────────────────

async def read_upload_file_bytes(file: UploadFile) -> tuple[bytes, str]:
    """
    Validate and read an uploaded file into memory without writing to disk.
    Returns (file_bytes, content_type).
    """
    if file.content_type not in settings.ALLOWED_FILE_TYPES:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"File type '{file.content_type}' is not allowed. Only PDF and DOCX are accepted.",
        )

    content = await file.read()
    if len(content) > settings.MAX_FILE_SIZE:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail="File exceeds maximum size of 10MB",
        )

    return content, file.content_type or "application/pdf"


def extract_text_from_pdf_bytes(file_bytes: bytes) -> str:
    """Extract plain text from PDF bytes using pdfplumber (no temp file)."""
    import io
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            pages = [page.extract_text() or "" for page in pdf.pages]
        return "\n".join(pages)
    except Exception as exc:
        raise RuntimeError(f"Failed to extract text from PDF bytes: {exc}") from exc


def extract_text_from_docx_bytes(file_bytes: bytes) -> str:
    """Extract plain text from DOCX bytes using python-docx (no temp file)."""
    import io
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.append(cell.text)
        return "\n".join(paragraphs)
    except Exception as exc:
        raise RuntimeError(f"Failed to extract text from DOCX bytes: {exc}") from exc


def extract_text_from_bytes(file_bytes: bytes, content_type: str) -> str:
    """Route to the correct BYTEA extractor based on content type."""
    if content_type == "application/pdf":
        return extract_text_from_pdf_bytes(file_bytes)
    elif content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return extract_text_from_docx_bytes(file_bytes)
    else:
        raise ValueError(f"Unsupported content type: {content_type}")
